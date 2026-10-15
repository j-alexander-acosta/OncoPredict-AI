from docx import Document
from model_service import ModelManager
import sys

try:
    manager = ModelManager()
    tabular_metrics = manager.get_tabular_metrics()
    image_metrics = manager.get_image_metrics()

    doc = Document()
    doc.add_heading('Análisis Comparativo - Métricas de Evaluación e Hiperparámetros', 0)

    # ==========================================
    # Tabla 1: Métricas de Modelos Tabulares (UCI)
    # ==========================================
    doc.add_heading('1. Métricas de Evaluación en Datos Clínicos Tabulares (15% Estratificado)', level=2)
    p1 = doc.add_paragraph('Evaluación rigurosa sobre 129 casos de prueba (121 negativos, 8 positivos) a partir de 32 factores de riesgo clínicos y epidemiológicos (UCI) sin fuga de información (leakage-free).')
    
    headers_tab = ["Algoritmo", "Dataset", "Accuracy", "Sensibilidad", "Especificidad", "Precisión", "F1-Score", "AUC-ROC", "Matriz (TP/FP/TN/FN)"]
    table_tab = doc.add_table(rows=1, cols=len(headers_tab))
    table_tab.style = 'Table Grid'
    for i, h in enumerate(headers_tab):
        table_tab.rows[0].cells[i].text = h

    for m in tabular_metrics:
        row = table_tab.add_row().cells
        row[0].text = str(m.get('model', ''))
        row[1].text = str(m.get('dataset', ''))
        row[2].text = str(m.get('accuracy_pct', ''))
        row[3].text = str(m.get('sensitivity_pct', ''))
        row[4].text = str(m.get('specificity_pct', ''))
        row[5].text = str(m.get('precision_pct', ''))
        row[6].text = str(m.get('f1_pct', ''))
        row[7].text = str(m.get('auc_roc_pct', ''))
        row[8].text = f"TP:{m.get('tp',0)} FP:{m.get('fp',0)} TN:{m.get('tn',0)} FN:{m.get('fn',0)}"

    doc.add_paragraph()

    # ==========================================
    # Tabla 2: Métricas de Modelos de Imágenes (CV)
    # ==========================================
    doc.add_heading('2. Métricas de Evaluación en Diagnóstico por Imagen (Computer Vision)', level=2)
    p2 = doc.add_paragraph('Evaluación multiclasificación de redes neuronales convolucionales profundas (CNN) sobre conjuntos de prueba de imágenes citológicas microscópicas (Herlev, SIPaKMeD, RIVA) y macroscópicas colposcópicas (Intel & MobileODT).')

    headers_img = ["Algoritmo CNN", "Dataset", "Accuracy", "Sensibilidad (Macro)", "Especificidad", "Precisión (Macro)", "F1-Score", "AUC-ROC"]
    table_img = doc.add_table(rows=1, cols=len(headers_img))
    table_img.style = 'Table Grid'
    for i, h in enumerate(headers_img):
        table_img.rows[0].cells[i].text = h

    for m in image_metrics:
        row = table_img.add_row().cells
        row[0].text = str(m.get('model', ''))
        row[1].text = str(m.get('dataset', ''))
        row[2].text = str(m.get('accuracy_pct', ''))
        row[3].text = str(m.get('sensitivity_pct', ''))
        row[4].text = str(m.get('specificity_pct', ''))
        row[5].text = str(m.get('precision_pct', ''))
        row[6].text = str(m.get('f1_pct', ''))
        row[7].text = str(m.get('auc_roc_pct', ''))

    doc.add_paragraph()

    # ==========================================
    # Tabla 3: Pipeline y Preprocesamiento
    # ==========================================
    doc.add_heading('3. Resumen de Configuración y Pipeline de Preprocesamiento (Datos UCI)', level=2)
    pipeline_data = [
        ["Fase del Pipeline", "Configuración / Parámetro Aplicado", "Descripción"],
        ["Prevención de Fuga (Data Leakage)", "drop(columns=['Hinselmann', 'Schiller', 'Citology'])", "Se eliminan variables de diagnóstico intermedio que sesgarían el modelo hacia la variable objetivo (Biopsy)."],
        ["División de Datos (Hold-out)", "test_size=0.15", "Se reserva un 15% del dataset exclusivamente para prueba (evaluación final)."],
        ["Estratificación", "stratify=y", "Garantiza que la proporción original de clases positivas y negativas se mantenga en entrenamiento y prueba."],
        ["Tratamiento de Nulos", "SimpleImputer(strategy='median')", "Los valores faltantes ('?') se reemplazan por la mediana de cada columna."],
        ["Escalado / Estandarización", "StandardScaler()", "Ajuste para que las variables numéricas tengan media 0 y varianza 1."],
        ["Balanceo de Clases", "SMOTE(random_state=42)", "Generación de muestras sintéticas para balancear las clases, aplicado únicamente al conjunto de entrenamiento."],
        ["Reproducibilidad", "random_state=42", "Semilla fija en todas las fases estocásticas para que los resultados sean deterministas y reproducibles."]
    ]

    table3 = doc.add_table(rows=1, cols=3)
    table3.style = 'Table Grid'
    for i in range(3):
        table3.rows[0].cells[i].text = pipeline_data[0][i]

    for r in pipeline_data[1:]:
        row_cells = table3.add_row().cells
        for i in range(3):
            row_cells[i].text = r[i]

    doc.add_paragraph()

    # ==========================================
    # Tabla 4: Hiperparámetros de Modelos Tabulares
    # ==========================================
    doc.add_heading('4. Resumen de Hiperparámetros por Modelo Tabular', level=2)
    hyperparam_data = [
        ["Modelo", "Hiperparámetros Clave", "Notas"],
        ["Regresión Logística", "max_iter=1000, random_state=42", "Se aumentó el límite de iteraciones (1000) para garantizar la convergencia del optimizador."],
        ["SVM (Vectores de Soporte)", "kernel='rbf', probability=True, random_state=42", "probability=True es necesario para poder calcular la métrica AUC-ROC en la evaluación."],
        ["Decision Tree", "random_state=42 (Sin límite de profundidad)", "Árbol de decisión estándar utilizando los parámetros por defecto de scikit-learn."],
        ["Random Forest", "n_estimators=100, random_state=42", "Ensamblado de 100 árboles de decisión para mayor robustez frente al sobreajuste."],
        ["XGBoost", "eval_metric='logloss', random_state=42", "Función de pérdida logarítmica, óptima para el problema de clasificación binaria abordado."]
    ]

    table4 = doc.add_table(rows=1, cols=3)
    table4.style = 'Table Grid'
    for i in range(3):
        table4.rows[0].cells[i].text = hyperparam_data[0][i]

    for r in hyperparam_data[1:]:
        row_cells = table4.add_row().cells
        for i in range(3):
            row_cells[i].text = r[i]

    output_path = '/Users/alexanderacosta/Documents/Proyectos/Análisis Comparativo/Resumen_Configuracion_e_Hiperparametros.docx'
    doc.save(output_path)
    print(f"Word generado exitosamente en: {output_path}")

except Exception as e:
    print(f"Error: {e}")
    sys.exit(1)
